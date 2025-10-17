#!/usr/bin/env python3
import os
import json
import glob
import pandas as pd
from typing import Dict, List
import hmac
import hashlib
import requests
from docx import Document


def norm_space(s: str) -> str:
    return " ".join((s or "").split()).strip()


def save_outputs(df: pd.DataFrame, base: str, out_dir: str) -> Dict[str,str]:
    os.makedirs(out_dir, exist_ok=True)
    csv_path = os.path.join(out_dir, base + '.csv')
    xlsx_path = os.path.join(out_dir, base + '.xlsx')
    docx_path = os.path.join(out_dir, base + '.docx')
    df.to_csv(csv_path, index=False)
    df.to_excel(xlsx_path, index=False)
    # DOCX
    doc = Document()
    table = doc.add_table(rows=len(df)+1, cols=len(df.columns))
    for j, col in enumerate(df.columns):
        table.cell(0,j).text = col
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, val in enumerate(row):
            table.cell(i,j).text = '' if pd.isna(val) else str(val)
    doc.save(docx_path)
    return { 'csv': csv_path, 'xlsx': xlsx_path, 'docx': docx_path }


def compute_signature(secret: str, payload: bytes) -> str:
    return hmac.new(secret.encode('utf-8'), payload, hashlib.sha256).hexdigest()


def upload_results(upload_url: str, token: str, doc_id: str, paths: Dict[str,str], summary: Dict) -> None:
    if not upload_url:
        print('[agg] RESULT_UPLOAD_URL empty; skipping upload to app')
        return
    try:
        files = {}
        for key in ['csv','xlsx','docx']:
            p = paths.get(key)
            if p and os.path.exists(p):
                files[key] = (os.path.basename(p), open(p, 'rb'), 'application/octet-stream')
        data = { 'doc_id': doc_id or '', 'summary': json.dumps(summary) }
        headers = {}
        if token:
            headers['Authorization'] = f"Bearer {token}"
        resp = requests.post(upload_url, files=files, data=data, headers=headers, timeout=120)
        print('[agg] Upload results status:', resp.status_code)
    except Exception as e:
        print('[agg] Upload results error:', repr(e))


def main():
    chunks_dir = os.getenv('CHUNKS_DIR', 'chunks')
    original_filename = os.getenv('ORIGINAL_FILENAME', 'document.pdf')
    callback_url = os.getenv('CALLBACK_URL', '')
    callback_secret = os.getenv('CALLBACK_HMAC_SECRET', '')
    result_upload_url = os.getenv('RESULT_UPLOAD_URL', '')
    result_upload_token = os.getenv('RESULT_UPLOAD_TOKEN', '')
    doc_id = os.getenv('DOC_ID', '')

    base = os.path.splitext(os.path.basename(original_filename))[0]
    pattern = os.path.join(chunks_dir, f"{base}-p*/*.csv")
    csv_files = sorted(glob.glob(pattern))
    if not csv_files:
        # also try flat files pattern
        csv_files = sorted(glob.glob(os.path.join(chunks_dir, f"**/{base}-p*.csv"), recursive=True))
    if not csv_files:
        raise SystemExit(f"No chunk CSV files found under {chunks_dir} for base {base}")

    frames: List[pd.DataFrame] = []
    for f in csv_files:
        try:
            df = pd.read_csv(f)
            frames.append(df)
            print(f"[agg] Loaded {f} -> {len(df)} rows")
        except Exception as e:
            print(f"[agg] Failed to read {f}: {e}")
    if not frames:
        raise SystemExit("No valid CSVs to aggregate")
    df_all = pd.concat(frames, ignore_index=True)
    # Normalize columns and deduplicate
    cols = ['surname','first_name','other_name','course_studied','faculty','grade','qualification_obtained','session']
    for c in cols:
        if c in df_all.columns:
            df_all[c] = df_all[c].astype(str).map(norm_space)
        else:
            df_all[c] = ''
    df_all = df_all[cols]
    before = len(df_all)
    df_all = df_all.drop_duplicates()
    after = len(df_all)
    print(f"[agg] Aggregated rows: {before} -> {after} (deduped)")

    out_dir = os.path.join('outputs')
    paths = save_outputs(df_all, base, out_dir)

    summary = {
        'status': 'success',
        'counts': { 'rows': int(after) },
        'chunks': len(frames),
        'files': paths,
        'filename': original_filename,
    }

    # Upload to app
    if not result_upload_token:
        print('[agg] RESULT_UPLOAD_TOKEN is empty; upload may be rejected by app')
    upload_results(result_upload_url, result_upload_token, doc_id, paths, summary)

    # Callback to app
    if callback_url:
        payload = json.dumps(summary).encode('utf-8')
        headers = { 'Content-Type': 'application/json' }
        if callback_secret:
            headers['X-Extractor-Signature'] = compute_signature(callback_secret, payload)
        try:
            cr = requests.post(callback_url, data=payload, headers=headers, timeout=60)
            print('[agg] Callback status:', cr.status_code)
        except Exception as e:
            print('[agg] Callback error:', repr(e))
    else:
        print('[agg] CALLBACK_URL not set; skipping callback')

    print(json.dumps(summary, indent=2))


if __name__ == '__main__':
    main()
